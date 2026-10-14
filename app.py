
from __future__ import annotations

from io import BytesIO
import re

import pandas as pd
import plotly.express as px
import streamlit as st


# =========================================================
# 1. PAGE SETUP AND DESIGN
# =========================================================

st.set_page_config(
    page_title="Student Result Intelligence",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
.stApp, [data-testid="stAppViewContainer"] {
    background: #080b14;
    color: #f8fafc;
}
[data-testid="stHeader"] {
    background: #080b14;
}
[data-testid="stSidebar"] {
    background: #11172b;
}
.block-container {
    max-width: 1500px;
    padding-top: 1.3rem;
    padding-bottom: 2rem;
}
h1, h2, h3, h4 {
    color: #c7d2fe !important;
}
.stApp p, .stApp label, .stApp li,
[data-testid="stWidgetLabel"] {
    color: #e2e8f0;
}
.hero {
    background: linear-gradient(120deg, #27245f, #4338ca, #155e75);
    padding: 26px;
    border-radius: 18px;
    margin-bottom: 22px;
}
.hero h1 {
    color: #ffffff !important;
    margin: 0;
}
.hero p {
    color: #e0e7ff !important;
    margin-bottom: 0;
}
div[data-testid="stMetric"] {
    background: #131b2e;
    border: 1px solid #34415e;
    padding: 15px;
    border-radius: 13px;
}
div[data-testid="stMetricLabel"] *,
div[data-testid="stMetricValue"] * {
    color: #ffffff !important;
}
.stButton button, .stDownloadButton button {
    background: #4338ca !important;
    color: #ffffff !important;
    border: 1px solid #6366f1 !important;
    border-radius: 9px;
    min-height: 40px;
}
.stButton button *, .stDownloadButton button * {
    color: #ffffff !important;
}
.stTextInput input, .stNumberInput input,
[data-baseweb="select"] > div {
    background: #111a2d !important;
    color: #ffffff !important;
}
[data-testid="stFileUploader"] {
    background: #11182b;
    border: 1px dashed #818cf8;
    border-radius: 12px;
    padding: 12px;
}
[data-testid="stDataFrame"] {
    border: 1px solid #34415e;
    border-radius: 10px;
}
hr {
    border-color: #34415e;
}
</style>

<div class="hero">
    <h1>🎓 Student Result Intelligence</h1>
    <p>Academic analytics · Result summary · SGPA/CGPA categories · Student profiles</p>
</div>
""", unsafe_allow_html=True)


# =========================================================
# 2. COLUMN NORMALIZATION
# =========================================================

ALIASES = {
    "seat_no": [
        "seat no", "seat number", "seat_no", "roll no",
        "roll number", "roll_no", "prn", "student id",
        "enrollment no", "enrolment no", "exam seat no",
    ],
    "student_name": [
        "student name", "student_name", "name",
        "candidate name", "name of student",
    ],
    "result": [
        "result", "status", "final result", "result status",
        "pass fail", "exam result",
    ],
    "sgpa": ["sgpa", "semester grade point average"],
    "cgpa": ["cgpa", "cumulative grade point average"],
    "total_marks": [
        "total marks", "total_marks", "grand total",
        "marks obtained", "total obtained",
    ],
    "percentage": [
        "percentage", "percent", "percentage marks",
        "aggregate percentage",
    ],
}

META_COLUMNS = {
    "seat_no", "student_name", "result", "sgpa", "cgpa",
    "total_marks", "percentage", "source_page",
    "semester", "institute", "course", "year",
    "teacher", "teacher_name", "class_category",
    "appeared", "maximum_marks",
}


def clean(value):
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_header(value):
    return re.sub(
        r"[^a-z0-9]+", "_", str(value).lower()
    ).strip("_")


def normalize_columns(df):
    df = df.copy()
    df.columns = [normalize_header(c) for c in df.columns]

    rename = {}

    for target, aliases in ALIASES.items():
        candidates = {normalize_header(x) for x in aliases}
        candidates.add(target)

        for col in df.columns:
            if col in candidates:
                rename[col] = target
                break

    df = df.rename(columns=rename)

    if "student_name" not in df.columns:
        candidates = [
            c for c in df.columns
            if "name" in c
            and "college" not in c
            and "institute" not in c
        ]
        if candidates:
            df = df.rename(
                columns={candidates[0]: "student_name"}
            )

    if "seat_no" not in df.columns:
        candidates = [
            c for c in df.columns
            if any(x in c for x in [
                "seat", "roll", "prn", "enrollment"
            ])
        ]
        if candidates:
            df = df.rename(
                columns={candidates[0]: "seat_no"}
            )

    if "total_marks" not in df.columns:
        candidates = [
            c for c in df.columns
            if "total" in c and "mark" in c
        ]
        if candidates:
            df = df.rename(
                columns={candidates[0]: "total_marks"}
            )

    return df


def numeric_series(series):
    return pd.to_numeric(
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("%", "", regex=False)
        .str.strip(),
        errors="coerce",
    )


# =========================================================
# 3. SUBJECT DETECTION
# =========================================================

def subject_code(column):
    text = str(column).upper()
    match = re.search(
        r"\b[A-Z]{2,}[A-Z0-9]*\d[A-Z0-9]*\b",
        text,
    )
    return match.group(0) if match else ""


def detect_subjects(df):
    subjects = []

    for col in list(df.columns):
        if col in META_COLUMNS:
            continue

        if any(
            word in col
            for word in [
                "teacher", "faculty", "institute", "college",
                "subject_name", "subject_code", "maximum",
                "max_marks", "source_page", "percentage",
                "sgpa", "cgpa",
            ]
        ):
            continue

        values = numeric_series(df[col])

        if len(df) > 0 and values.notna().sum() >= max(
            2, int(len(df) * 0.4)
        ):
            df[col] = values
            subjects.append(col)

    return subjects


def extract_subject_names(source_text, subjects):
    mapping = {}

    if not source_text:
        return mapping

    pattern = re.compile(
        r"([A-Z]{2,}[A-Z0-9]*\d[A-Z0-9]*)\s*[:：-]\s*",
        re.I,
    )

    for line in source_text.splitlines():
        line = re.sub(r"\s+", " ", line).strip()
        matches = list(pattern.finditer(line))

        for i, match in enumerate(matches):
            end = (
                matches[i + 1].start()
                if i + 1 < len(matches)
                else len(line)
            )
            name = line[match.end():end].strip(" :-")

            if name and len(name) > 2:
                mapping[match.group(1).upper()] = name.title()

    result = {}

    for col in subjects:
        code = subject_code(col)
        if code in mapping:
            result[col] = mapping[code]

    return result


def display_subject(column, subject_names):
    if column in subject_names:
        return subject_names[column]

    code = subject_code(column)

    if code:
        return f"{code} — Subject name not detected"

    return str(column).replace("_", " ").title()


# =========================================================
# 4. DATA PREPARATION
# =========================================================

def calculate_percentage(df, original_columns, source_text):
    df = df.copy()

    if "percentage" not in df.columns:
        df["percentage"] = pd.NA
    else:
        df["percentage"] = numeric_series(df["percentage"])

    maximum = None

    for column in original_columns:
        label = str(column)

        match = re.search(
            r"(?:/|\(|out of\s*)(\d{2,4})",
            label,
            re.I,
        )

        if (
            match
            and re.search(
                r"total.*mark|mark.*total|grand.*total",
                label,
                re.I,
            )
        ):
            maximum = float(match.group(1))
            break

    if maximum is None and source_text:
        match = re.search(
            r"Total\s*Marks?\s*\(\s*(\d{2,4})\s*\)",
            source_text,
            re.I,
        )
        if match:
            maximum = float(match.group(1))

    if (
        df["percentage"].notna().sum() == 0
        and maximum
        and "total_marks" in df.columns
    ):
        df["percentage"] = (
            numeric_series(df["total_marks"]) / maximum * 100
        ).clip(0, 100).round(2)

    # This is only an estimate; it is not an official percentage.
    if "sgpa" in df.columns:
        df["sgpa_percentage_estimate"] = (
            (numeric_series(df["sgpa"]) - 0.75) * 10
        ).clip(0, 100).round(2)
    else:
        df["sgpa_percentage_estimate"] = pd.NA

    return df


def prepare_data(raw, source_text=""):
    if raw is None or raw.empty:
        raise ValueError("The uploaded file contains no readable rows.")

    original_columns = list(raw.columns)
    df = normalize_columns(raw)

    if "student_name" not in df.columns:
        raise ValueError(
            "Student Name column was not detected. "
            "Use a structured result table with a student-name column."
        )

    if "seat_no" not in df.columns:
        df["seat_no"] = [
            f"ROW-{i + 1:04d}" for i in range(len(df))
        ]

    result_was_provided = "result" in df.columns

    if not result_was_provided:
        df["result"] = ""

    df["student_name"] = df["student_name"].map(clean)
    df["seat_no"] = (
        df["seat_no"].map(clean)
        .str.replace(r"\.0$", "", regex=True)
    )
    df["result"] = df["result"].map(clean).str.upper()

    for col in ["sgpa", "cgpa", "total_marks", "percentage"]:
        if col in df.columns:
            df[col] = numeric_series(df[col])

    subjects = detect_subjects(df)

    if "total_marks" not in df.columns and subjects:
        df["total_marks"] = df[subjects].sum(
            axis=1, min_count=1
        )

    df = calculate_percentage(df, original_columns, source_text)
    df.attrs["result_was_provided"] = result_was_provided

    return df, subjects


# =========================================================
# 5. PDF, CSV AND EXCEL IMPORT
# =========================================================

def read_pdf(upload):
    try:
        import pdfplumber
    except ImportError as exc:
        raise ValueError(
            "PDF package missing. Run: "
            "python -m pip install pdfplumber"
        ) from exc

    upload.seek(0)
    extracted_tables = []
    text_parts = []

    with pdfplumber.open(upload) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")

            try:
                tables = page.extract_tables() or []
            except Exception:
                tables = []

            for table in tables:
                if not table or len(table) < 2:
                    continue

                rows = [
                    [clean(cell) for cell in row]
                    for row in table if row
                ]

                header_index = None

                for i, row in enumerate(rows[:5]):
                    line = " ".join(row).lower()

                    if (
                        "name" in line
                        and any(
                            word in line
                            for word in [
                                "seat", "roll", "prn",
                                "result", "marks",
                            ]
                        )
                    ):
                        header_index = i
                        break

                if header_index is None:
                    continue

                headers = [
                    clean(c) or f"column_{i + 1}"
                    for i, c in enumerate(rows[header_index])
                ]

                for row in rows[header_index + 1:]:
                    row = row + [""] * max(
                        0, len(headers) - len(row)
                    )
                    row = row[:len(headers)]

                    if not any(row):
                        continue

                    extracted_tables.append(dict(zip(headers, row)))

    full_text = "\n".join(text_parts)

    if extracted_tables:
        raw = pd.DataFrame(extracted_tables).drop_duplicates()
        return raw.reset_index(drop=True), "PDF table extraction", full_text

    pattern = re.compile(
        r"(?P<seat>\b\d{5,15}\b)\s+"
        r"(?P<name>[A-Za-z][A-Za-z .,'()-]{2,}?)\s+"
        r"(?P<result>PASS|FAIL|ATKT|ABSENT|PROMOTED)\s+"
        r"(?P<sgpa>\d(?:\.\d{1,2})?)\s+"
        r"(?P<total>\d{1,4})\b",
        re.I,
    )

    records = []

    for match in pattern.finditer(full_text):
        records.append({
            "Seat No": match.group("seat"),
            "Student Name": match.group("name").strip(),
            "Result": match.group("result").upper(),
            "SGPA": match.group("sgpa"),
            "Total Marks": match.group("total"),
        })

    if records:
        return pd.DataFrame(records), "PDF text fallback", full_text

    raise ValueError(
        "The PDF result table could not be extracted reliably. "
        "It may be scanned or have a different layout. "
        "Try a searchable PDF or export the result table to Excel/CSV. "
        "The app will not guess missing marks or statuses."
    )


def load_upload(upload):
    extension = upload.name.lower().rsplit(".", 1)[-1]

    if extension == "csv":
        return pd.read_csv(upload), "CSV import", ""

    if extension in ["xlsx", "xls"]:
        xls = pd.ExcelFile(upload)
        candidates = []

        for sheet in xls.sheet_names:
            raw = pd.read_excel(xls, sheet_name=sheet)

            try:
                normalized, subjects = prepare_data(raw)
                candidates.append(
                    (len(normalized), len(subjects), raw, sheet)
                )
            except Exception:
                continue

        if not candidates:
            raise ValueError(
                "No Excel sheet with a detectable student-name column was found."
            )

        _, _, raw, sheet = max(
            candidates, key=lambda item: (item[0], item[1])
        )

        return raw, f"Excel import · {sheet}", ""

    if extension == "pdf":
        return read_pdf(upload)

    raise ValueError("Supported formats: PDF, CSV, XLSX and XLS.")


# =========================================================
# 6. EXPORT HELPERS
# =========================================================

def excel_bytes(df):
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(
            writer, index=False, sheet_name="Student Results"
        )

    return output.getvalue()


def word_report(df, subjects, subject_names, source_name, semester):
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError(
            "Word export needs python-docx. Run: "
            "python -m pip install python-docx"
        ) from exc

    doc = Document()
    doc.add_heading("Student Result Intelligence Report", 0)
    doc.add_paragraph(f"Source: {source_name}")
    doc.add_paragraph(f"Semester / Exam: {semester}")
    doc.add_paragraph(f"Students analysed: {len(df)}")

    if df["percentage"].notna().any():
        doc.add_paragraph(
            f"Average percentage: {df['percentage'].mean():.2f}%"
        )
    else:
        doc.add_paragraph(
            "Official percentage is not available in the uploaded data."
        )

    if "sgpa" in df.columns and df["sgpa"].notna().any():
        doc.add_paragraph(
            f"Average SGPA: {df['sgpa'].mean():.2f}"
        )

    ranking_column = next(
        (
            col for col in ["total_marks", "percentage", "sgpa"]
            if col in df.columns
            and numeric_series(df[col]).notna().any()
        ),
        None,
    )

    if ranking_column:
        doc.add_heading("Top 10 Students", level=1)
        ranked = df.copy()
        ranked[ranking_column] = numeric_series(ranked[ranking_column])
        ranked = ranked.sort_values(
            ranking_column, ascending=False, na_position="last"
        ).head(10)

        for _, row in ranked.iterrows():
            doc.add_paragraph(
                f"{row['student_name']} | {row['seat_no']} | "
                f"{ranking_column}: {row[ranking_column]}"
            )

    doc.add_heading("Subject Intelligence", level=1)

    for subject in subjects:
        values = numeric_series(df[subject]).dropna()

        if values.empty:
            continue

        doc.add_paragraph(
            f"{subject_names.get(subject, subject)} | "
            f"Average: {values.mean():.2f} | "
            f"Highest: {values.max():g} | "
            f"Lowest: {values.min():g}"
        )

    doc.add_heading("Student Register", level=1)

    columns = [
        c for c in [
            "seat_no", "student_name", "result", "sgpa", "cgpa",
            "total_marks", "percentage",
        ] + subjects if c in df.columns
    ]

    if columns:
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"

        for i, col in enumerate(columns):
            table.rows[0].cells[i].text = col.replace("_", " ").title()

        for _, row in df.iterrows():
            cells = table.add_row().cells

            for i, col in enumerate(columns):
                value = row.get(col, "")
                cells[i].text = (
                    "" if pd.isna(value) else str(value)[:100]
                )

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


# =========================================================
# 7. UPLOAD DATA
# =========================================================

with st.sidebar:
    st.markdown("## 📂 Import Result Data")

    uploaded = st.file_uploader(
        "Upload college result",
        type=["pdf", "csv", "xlsx", "xls"],
    )

    semester = st.text_input(
        "Semester / Examination",
        "Semester Result",
    )

    st.caption(
        "After importing, verify the extracted data against the official result."
    )

if uploaded is None:
    st.info(
        "Upload your college result PDF, CSV or Excel file from the sidebar."
    )
    st.stop()

try:
    raw_df, import_method, source_text = load_upload(uploaded)
    df, subjects = prepare_data(raw_df, source_text)
    subject_names = extract_subject_names(source_text, subjects)

except Exception as exc:
    st.error(f"Import failed: {exc}")
    st.info(
        "If the PDF is scanned, OCR may be required. "
        "Try a searchable PDF or an Excel/CSV export."
    )
    st.stop()

source_name = uploaded.name
result_was_provided = bool(df.attrs.get("result_was_provided", False))

st.caption(
    f"File: {source_name} · Import: {import_method} · "
    f"Students: {len(df)} · Subjects detected: {len(subjects)}"
)


# =========================================================
# 8. RESULT STATUS LOGIC
# =========================================================

status = df["result"].fillna("").astype(str).str.upper().str.strip()

atkt_mask = status.str.contains(
    r"\bATKT\b|\bBACKLOG\b",
    regex=True,
    na=False,
)

fail_mask = status.str.contains(
    r"\bFAIL(?:ED)?\b",
    regex=True,
    na=False,
) & ~atkt_mask

absent_mask = status.str.contains(
    r"\bABSENT\b|\bAB\b",
    regex=True,
    na=False,
)

pending_mask = status.str.contains(
    r"PENDING|WITHHELD|RESERVED|AWAITED|NOT DECLARED",
    regex=True,
    na=False,
)

promoted_mask = status.str.contains(
    r"\bPROMOTED\b",
    regex=True,
    na=False,
) & ~atkt_mask & ~fail_mask & ~absent_mask & ~pending_mask

pass_mask = status.str.contains(
    r"\bPASS(?:ED)?\b",
    regex=True,
    na=False,
) & ~atkt_mask & ~fail_mask & ~absent_mask & ~pending_mask

all_clear_mask = (
    (pass_mask | promoted_mask)
    & ~atkt_mask
    & ~fail_mask
    & ~absent_mask
    & ~pending_mask
)

passed_promoted_mask = all_clear_mask | atkt_mask

if result_was_provided:
    known_appeared_status = all_clear_mask | atkt_mask | fail_mask
    appeared_count = int(known_appeared_status.sum())
else:
    appeared_count = None

total_students = len(df)
all_clear_count = int(all_clear_mask.sum())
atkt_count = int(atkt_mask.sum())
fail_count = int(fail_mask.sum())
pending_count = int(pending_mask.sum())
absent_count = int(absent_mask.sum())
passed_promoted_count = int(passed_promoted_mask.sum())

sgpa_values = (
    numeric_series(df["sgpa"])
    if "sgpa" in df.columns
    else pd.Series(float("nan"), index=df.index)
)

cgpa_values = (
    numeric_series(df["cgpa"])
    if "cgpa" in df.columns
    else pd.Series(float("nan"), index=df.index)
)

avg_percentage = numeric_series(df["percentage"]).mean()
avg_sgpa = sgpa_values.mean()

avg_total = (
    numeric_series(df["total_marks"]).mean()
    if "total_marks" in df.columns
    else float("nan")
)


# =========================================================
# 9. NAVIGATION
# =========================================================

pages = [
    "Academic Overview",
    "Result Summary & Class Category",
    "Top Performers",
    "Individual Student Analysis",
    "Student Risk",
    "Student Explorer",
    "Import Results",
    "Export Centre",
]

selected = st.radio(
    "Navigation",
    pages,
    horizontal=True,
    label_visibility="visible",
)


# =========================================================
# 10. ACADEMIC OVERVIEW
# =========================================================

if selected == "Academic Overview":
    st.header("📊 Academic Overview")

    a, b, c, d = st.columns(4)

    a.metric("Total Students", total_students)
    b.metric(
        "Students Appeared",
        appeared_count if appeared_count is not None else "N/A",
    )
    c.metric("Passed / Promoted", passed_promoted_count)
    d.metric(
        "Average Percentage",
        f"{avg_percentage:.2f}%"
        if pd.notna(avg_percentage) else "N/A",
    )

    e, f, g = st.columns(3)

    e.metric(
        "Average SGPA",
        f"{avg_sgpa:.2f}" if pd.notna(avg_sgpa) else "N/A",
    )
    f.metric(
        "Average Total Marks",
        f"{avg_total:.2f}" if pd.notna(avg_total) else "N/A",
    )
    g.metric("ATKT Students", atkt_count)

    status_chart = pd.DataFrame([
        {"Status": "All Clear", "Students": all_clear_count},
        {"Status": "ATKT", "Students": atkt_count},
        {"Status": "Fail", "Students": fail_count},
        {"Status": "Absent", "Students": absent_count},
        {"Status": "Result Pending", "Students": pending_count},
    ])

    status_chart = status_chart[status_chart["Students"] > 0]

    if not status_chart.empty:
        st.plotly_chart(
            px.pie(
                status_chart,
                names="Status",
                values="Students",
                title="Result Status Distribution",
                hole=0.42,
            ),
            use_container_width=True,
        )

    if numeric_series(df["percentage"]).notna().any():
        chart_df = df.copy()
        chart_df["percentage"] = numeric_series(chart_df["percentage"])

        st.plotly_chart(
            px.histogram(
                chart_df.dropna(subset=["percentage"]),
                x="percentage",
                nbins=15,
                title="Student Percentage Distribution",
            ),
            use_container_width=True,
        )

    if subjects:
        averages = pd.DataFrame({
            "Subject": [
                display_subject(s, subject_names) for s in subjects
            ],
            "Average Marks": [
                numeric_series(df[s]).mean() for s in subjects
            ],
        }).dropna()

        st.plotly_chart(
            px.bar(
                averages,
                x="Subject",
                y="Average Marks",
                title="Subject-wise Average Marks",
                color="Average Marks",
                color_continuous_scale="Purples",
            ),
            use_container_width=True,
        )

    st.subheader("Result Data Preview")
    st.dataframe(df.head(30), use_container_width=True, hide_index=True)


# =========================================================
# 11. RESULT SUMMARY & CLASS CATEGORY
# =========================================================

elif selected == "Result Summary & Class Category":
    st.header("📋 Result Summary & Class Category")

    st.caption(
        "Counts and percentages are calculated from the uploaded result. "
        "Missing official statuses are not guessed."
    )

    def count_percent(count, denominator):
        if denominator == 0:
            return "N/A"
        return f"{count / denominator * 100:.2f}%"

    # -----------------------------------------------------
    # A: TOTAL RESULT
    # -----------------------------------------------------

    st.subheader("A: Total Result")

    result_summary = pd.DataFrame([
        {
            "Particulars": "Total Number of Students",
            "No.": total_students,
            "%": "100.00%" if total_students else "N/A",
        },
        {
            "Particulars": "Total Number of Students Appeared",
            "No.": appeared_count if appeared_count is not None else "N/A",
            "%": (
                count_percent(appeared_count, total_students)
                if appeared_count is not None else "N/A"
            ),
        },
        {
            "Particulars": "Students Passed / Promoted",
            "No.": passed_promoted_count,
            "%": count_percent(passed_promoted_count, total_students),
        },
        {
            "Particulars": "Students All Clear",
            "No.": all_clear_count,
            "%": count_percent(all_clear_count, total_students),
        },
        {
            "Particulars": "Students with ATKT",
            "No.": atkt_count,
            "%": count_percent(atkt_count, total_students),
        },
        {
            "Particulars": "Students Failed",
            "No.": fail_count,
            "%": count_percent(fail_count, total_students),
        },
        {
            "Particulars": "Result Pending",
            "No.": pending_count,
            "%": count_percent(pending_count, total_students),
        },
        {
            "Particulars": "Students Absent",
            "No.": absent_count,
            "%": count_percent(absent_count, total_students),
        },
    ])

    st.dataframe(
        result_summary,
        use_container_width=True,
        hide_index=True,
    )

    m1, m2, m3, m4 = st.columns(4)

    m1.metric("All Clear", all_clear_count)
    m2.metric("ATKT", atkt_count)
    m3.metric("Fail", fail_count)
    m4.metric("Pending", pending_count)

    result_chart = pd.DataFrame([
        {"Category": "All Clear", "Students": all_clear_count},
        {"Category": "ATKT", "Students": atkt_count},
        {"Category": "Fail", "Students": fail_count},
        {"Category": "Result Pending", "Students": pending_count},
        {"Category": "Absent", "Students": absent_count},
    ])

    result_chart = result_chart[result_chart["Students"] > 0]

    if not result_chart.empty:
        st.plotly_chart(
            px.pie(
                result_chart,
                names="Category",
                values="Students",
                title="Result Status Distribution",
                hole=0.43,
            ),
            use_container_width=True,
        )

    if not result_was_provided:
        st.warning(
            "No detectable Result/Status column was found. "
            "Appeared, All Clear, ATKT and Fail cannot be determined reliably."
        )

    st.divider()

    # -----------------------------------------------------
    # B: SGPA / CGPA CATEGORIES + EXACT GRADE STUDENT LISTS
    # -----------------------------------------------------

    st.subheader("B: Class Category")

    category_method = st.radio(
        "Select category method",
        ["SGPA / CGPA", "Percentage-based class"],
        horizontal=True,
    )

    if category_method == "SGPA / CGPA":
        available_grade_columns = []

        if sgpa_values.notna().any():
            available_grade_columns.append("sgpa")

        if cgpa_values.notna().any():
            available_grade_columns.append("cgpa")

        if not available_grade_columns:
            st.warning(
                "SGPA/CGPA was not detected. Check the uploaded column headers "
                "or upload a result file containing SGPA/CGPA."
            )
        else:
            grade_column = st.selectbox(
                "Choose grade field",
                available_grade_columns,
                format_func=lambda col: col.upper(),
            )

            grade_values = numeric_series(df[grade_column])

            if result_was_provided:
                grade_eligible = (
                    passed_promoted_mask & grade_values.notna()
                )
                grade_basis = (
                    "Passed/Promoted records with numeric grades"
                )
            else:
                grade_eligible = grade_values.notna()
                grade_basis = (
                    "All records with numeric grades; official status unavailable"
                )

            eligible_grades = grade_values[grade_eligible]

            # RANGES: each grade belongs to one range only.
            grade_categories = [
                ("> 9", eligible_grades > 9),
                ("8 to 9", (eligible_grades >= 8) & (eligible_grades <= 9)),
                ("7 to <8", (eligible_grades >= 7) & (eligible_grades < 8)),
                ("6 to <7", (eligible_grades >= 6) & (eligible_grades < 7)),
                ("5 to <6", (eligible_grades >= 5) & (eligible_grades < 6)),
                ("Below 5", eligible_grades < 5),
            ]

            grade_rows = []

            for label, condition in grade_categories:
                count = int(condition.sum())
                grade_rows.append({
                    "SGPA / CGPA Range": label,
                    "No. of Students": count,
                    "% of Students with Grade": count_percent(
                        count, len(eligible_grades)
                    ),
                })

            grade_table = pd.DataFrame(grade_rows)

            st.caption(f"Basis: {grade_basis}")

            st.dataframe(
                grade_table,
                use_container_width=True,
                hide_index=True,
            )

            st.metric(
                "Students with Numeric Grade",
                len(eligible_grades),
            )

            if not grade_table.empty and grade_table["No. of Students"].sum() > 0:
                st.plotly_chart(
                    px.bar(
                        grade_table,
                        x="SGPA / CGPA Range",
                        y="No. of Students",
                        title=f"Student Count by {grade_column.upper()} Range",
                        text="No. of Students",
                    ),
                    use_container_width=True,
                )

                st.plotly_chart(
                    px.pie(
                        grade_table[grade_table["No. of Students"] > 0],
                        names="SGPA / CGPA Range",
                        values="No. of Students",
                        title=f"{grade_column.upper()} Range Distribution",
                        hole=0.42,
                    ),
                    use_container_width=True,
                )

            st.download_button(
                "Download SGPA / CGPA Category CSV",
                grade_table.to_csv(index=False).encode("utf-8-sig"),
                "sgpa_cgpa_category.csv",
                "text/csv",
            )

            # -------------------------------------------------
            # NEW FEATURE: EXACT SGPA/CGPA VALUES WITH NAMES
            # -------------------------------------------------

            st.divider()
            st.subheader(
                f"👥 Students with Exact {grade_column.upper()} Values"
            )

            st.write(
                "Here you can see the exact grade values separately, "
                "including each student's name and seat number."
            )

            exact_value_options = [9.0, 8.0, 7.0, 6.0, 5.0]

            exact_value_labels = [
                f"{value:.2f}" for value in exact_value_options
            ]

            chosen_exact_label = st.selectbox(
                f"Choose exact {grade_column.upper()} value",
                exact_value_labels,
                index=1,
                key="exact_grade_value",
            )

            chosen_exact_value = float(chosen_exact_label)

            # Compare rounded values to avoid float precision issues.
            rounded_grades = grade_values.round(2)

            exact_mask = (
                rounded_grades.eq(chosen_exact_value)
                & grade_eligible
            )

            exact_students = df.loc[exact_mask].copy()

            st.metric(
                f"Students with {grade_column.upper()} = "
                f"{chosen_exact_value:.2f}",
                len(exact_students),
            )

            if exact_students.empty:
                st.info(
                    f"No students with exact {grade_column.upper()} "
                    f"= {chosen_exact_value:.2f} were found in the uploaded data."
                )
            else:
                name_columns = [
                    col for col in [
                        "seat_no", "student_name", "result",
                        "sgpa", "cgpa", "total_marks", "percentage",
                    ]
                    if col in exact_students.columns
                ]

                exact_students = exact_students.sort_values(
                    by="student_name",
                    key=lambda s: s.astype(str).str.lower(),
                )

                st.dataframe(
                    exact_students[name_columns],
                    use_container_width=True,
                    hide_index=True,
                )

                st.download_button(
                    f"Download {grade_column.upper()} "
                    f"{chosen_exact_value:.2f} Students CSV",
                    exact_students[name_columns]
                    .to_csv(index=False)
                    .encode("utf-8-sig"),
                    f"{grade_column}_{chosen_exact_value:.2f}_students.csv",
                    "text/csv",
                )

            # Quick counts and named lists for exactly 8, 7 and 6.
            st.subheader(
                f"Quick View: Exact {grade_column.upper()} 8, 7 and 6"
            )

            exact_quick_rows = []

            for exact_value in [8.0, 7.0, 6.0]:
                mask = (
                    rounded_grades.eq(exact_value)
                    & grade_eligible
                )

                students_for_value = df.loc[mask]

                exact_quick_rows.append({
                    f"Exact {grade_column.upper()}": f"{exact_value:.2f}",
                    "Number of Students": len(students_for_value),
                    "Student Names": (
                        ", ".join(
                            students_for_value["student_name"]
                            .astype(str).tolist()
                        )
                        if not students_for_value.empty
                        else "No students found"
                    ),
                })

            exact_quick_table = pd.DataFrame(exact_quick_rows)

            st.dataframe(
                exact_quick_table,
                use_container_width=True,
                hide_index=True,
            )

            st.download_button(
                "Download Exact Grade 8, 7 and 6 Student Lists",
                exact_quick_table.to_csv(index=False).encode("utf-8-sig"),
                "exact_grade_8_7_6_students.csv",
                "text/csv",
            )

            st.caption(
                "Exact 8.00 means exactly 8.00, not the full 8-to-9 range. "
                "Use the range table above for all grades between 8 and 9."
            )

            missing_grade_count = int(grade_values.isna().sum())

            if missing_grade_count:
                st.info(
                    f"{missing_grade_count} record(s) have no numeric "
                    f"{grade_column.upper()} value. They are not assigned to a grade."
                )

    # -----------------------------------------------------
    # ALTERNATIVE: DISTINCTION / FIRST CLASS / ETC.
    # -----------------------------------------------------

    else:
        st.info(
            "University class thresholds can differ. These thresholds are "
            "editable; confirm them against your college's official rules."
        )

        c1, c2, c3 = st.columns(3)

        with c1:
            distinction_min = st.number_input(
                "Distinction starts at (%)",
                min_value=0.0,
                max_value=100.0,
                value=75.0,
                step=1.0,
            )

        with c2:
            first_min = st.number_input(
                "First Class starts at (%)",
                min_value=0.0,
                max_value=100.0,
                value=60.0,
                step=1.0,
            )

        with c3:
            second_min = st.number_input(
                "Second Class starts at (%)",
                min_value=0.0,
                max_value=100.0,
                value=45.0,
                step=1.0,
            )

        if not (
            distinction_min >= first_min >= second_min
        ):
            st.error(
                "Set thresholds in descending order: "
                "Distinction, First Class, Second Class."
            )

        elif not numeric_series(df["percentage"]).notna().any():
            st.warning(
                "Official percentage is not available. "
                "Class categories cannot be calculated accurately."
            )

        else:
            percentages = numeric_series(df["percentage"])

            if result_was_provided:
                eligible_pass = all_clear_mask
                eligible_atkt = atkt_mask
            else:
                eligible_pass = pd.Series(False, index=df.index)
                eligible_atkt = pd.Series(False, index=df.index)

            distinction_mask = (
                eligible_pass
                & percentages.notna()
                & (percentages >= distinction_min)
            )

            first_mask = (
                eligible_pass
                & percentages.notna()
                & (percentages >= first_min)
                & (percentages < distinction_min)
            )

            second_mask = (
                eligible_pass
                & percentages.notna()
                & (percentages >= second_min)
                & (percentages < first_min)
            )

            pass_class_mask = (
                eligible_pass
                & percentages.notna()
                & (percentages < second_min)
            )

            class_table = pd.DataFrame([
                {
                    "Class Category": "Distinction",
                    "No. of Students": int(distinction_mask.sum()),
                },
                {
                    "Class Category": "First Class",
                    "No. of Students": int(first_mask.sum()),
                },
                {
                    "Class Category": "Second Class",
                    "No. of Students": int(second_mask.sum()),
                },
                {
                    "Class Category": "Pass Class",
                    "No. of Students": int(pass_class_mask.sum()),
                },
                {
                    "Class Category": "ATKT",
                    "No. of Students": int(eligible_atkt.sum()),
                },
                {
                    "Class Category": "Fail",
                    "No. of Students": fail_count,
                },
            ])

            st.dataframe(
                class_table,
                use_container_width=True,
                hide_index=True,
            )

            if class_table["No. of Students"].sum() > 0:
                st.plotly_chart(
                    px.bar(
                        class_table,
                        x="Class Category",
                        y="No. of Students",
                        title="Class Category Distribution",
                        text="No. of Students",
                    ),
                    use_container_width=True,
                )

            if not result_was_provided:
                st.warning(
                    "Official PASS/ATKT status is missing. "
                    "The app cannot reliably determine official class categories."
                )

            st.download_button(
                "Download Class Category CSV",
                class_table.to_csv(index=False).encode("utf-8-sig"),
                "class_category_summary.csv",
                "text/csv",
            )

    st.divider()

    st.download_button(
        "Download Total Result Summary CSV",
        result_summary.to_csv(index=False).encode("utf-8-sig"),
        "total_result_summary.csv",
        "text/csv",
    )


# =========================================================
# 12. TOP PERFORMERS
# =========================================================

elif selected == "Top Performers":
    st.header("🏆 Top Performers")

    score_options = [
        col for col in ["total_marks", "percentage", "sgpa", "cgpa"]
        if col in df.columns and numeric_series(df[col]).notna().any()
    ]

    if not score_options:
        st.warning("No numeric ranking field was detected.")
    else:
        score_column = st.selectbox(
            "Rank students using",
            score_options,
            format_func=lambda col: col.replace("_", " ").title(),
        )

        ranked = df.copy()
        ranked[score_column] = numeric_series(ranked[score_column])
        ranked = ranked.dropna(subset=[score_column]).sort_values(
            score_column, ascending=False
        )

        top = ranked.head(10).copy()
        top.insert(0, "Rank", range(1, len(top) + 1))

        st.subheader("Top 10 Students")
        st.dataframe(top, use_container_width=True, hide_index=True)

        st.subheader("Lowest 10 Students")
        st.dataframe(
            ranked.sort_values(score_column).head(10),
            use_container_width=True,
            hide_index=True,
        )

        if not ranked.empty:
            winner = ranked.iloc[0]
            st.success(
                f"Top student: {winner['student_name']} "
                f"({winner['seat_no']}) — {winner[score_column]}"
            )

        st.download_button(
            "Download Ranking CSV",
            ranked.to_csv(index=False).encode("utf-8-sig"),
            "student_ranking.csv",
            "text/csv",
        )


# =========================================================
# 13. INDIVIDUAL STUDENT ANALYSIS
# =========================================================

elif selected == "Individual Student Analysis":
    st.header("🔎 Individual Student Analysis")

    query = st.text_input(
        "Search by Student Name or Seat Number",
        placeholder="Enter name or seat number",
        key="individual_search",
    )

    if query.strip():
        matches = df[
            df["student_name"].astype(str).str.contains(
                query, case=False, regex=False, na=False
            )
            | df["seat_no"].astype(str).str.contains(
                query, case=False, regex=False, na=False
            )
        ]
    else:
        matches = df.iloc[0:0]

    if matches.empty:
        st.info("Enter a name or seat number to view the student profile.")
    else:
        chosen_index = st.selectbox(
            "Select Student",
            matches.index.tolist(),
            format_func=lambda i: (
                f"{df.loc[i, 'student_name']} · {df.loc[i, 'seat_no']}"
            ),
            key="individual_select",
        )

        student = df.loc[chosen_index]

        st.subheader("👤 Student Profile")
        st.write(f"**Name:** {student['student_name']}")
        st.write(f"**Seat Number:** {student['seat_no']}")
        st.write(
            f"**Official Result:** {student.get('result', 'Not provided')}"
        )

        total = student.get("total_marks", pd.NA)
        percentage = student.get("percentage", pd.NA)
        sgpa = student.get("sgpa", pd.NA)
        estimate = student.get("sgpa_percentage_estimate", pd.NA)

        a, b, c, d = st.columns(4)

        a.metric(
            "Total Marks",
            f"{total:g}" if pd.notna(total) else "N/A",
        )
        b.metric(
            "Official Percentage",
            f"{percentage:.2f}%"
            if pd.notna(percentage) else "N/A",
        )
        c.metric(
            "SGPA",
            f"{sgpa:.2f}" if pd.notna(sgpa) else "N/A",
        )
        d.metric(
            "SGPA Percentage Estimate",
            f"{estimate:.2f}%" if pd.notna(estimate) else "N/A",
        )

        if subjects:
            subject_rows = []

            for subject in subjects:
                mark = pd.to_numeric(
                    pd.Series([student.get(subject, pd.NA)]),
                    errors="coerce",
                ).iloc[0]

                population = numeric_series(df[subject]).dropna()

                percentile = (
                    (population <= mark).sum() / len(population) * 100
                    if pd.notna(mark) and len(population)
                    else float("nan")
                )

                subject_rows.append({
                    "Subject": display_subject(subject, subject_names),
                    "Marks": mark,
                    "Percentile": percentile,
                })

            student_subjects = pd.DataFrame(subject_rows)

            st.subheader("Subject-wise Marks and Percentile")
            st.dataframe(
                student_subjects,
                use_container_width=True,
                hide_index=True,
            )

            left, right = st.columns(2)

            with left:
                st.plotly_chart(
                    px.bar(
                        student_subjects.dropna(subset=["Marks"]),
                        x="Subject",
                        y="Marks",
                        title="Subject-wise Marks",
                    ),
                    use_container_width=True,
                )

            with right:
                pie_df = student_subjects.dropna(subset=["Marks"])
                pie_df = pie_df[pie_df["Marks"] > 0]

                if not pie_df.empty:
                    st.plotly_chart(
                        px.pie(
                            pie_df,
                            names="Subject",
                            values="Marks",
                            title="Share of Obtained Subject Marks",
                            hole=0.42,
                        ),
                        use_container_width=True,
                    )

            st.plotly_chart(
                px.bar(
                    student_subjects.dropna(subset=["Percentile"]),
                    x="Subject",
                    y="Percentile",
                    title="Subject-wise Percentile Rank",
                    range_y=[0, 100],
                ),
                use_container_width=True,
            )

            st.caption(
                "Percentile is calculated against students in the uploaded file; "
                "it is not an official university percentile."
            )

        with st.expander("All Available Student Details"):
            st.dataframe(
                pd.DataFrame({
                    "Field": df.columns,
                    "Value": [
                        "" if pd.isna(student.get(col))
                        else str(student.get(col))
                        for col in df.columns
                    ],
                }),
                use_container_width=True,
                hide_index=True,
            )

        st.download_button(
            "Download Individual Student CSV",
            pd.DataFrame([student]).to_csv(index=False).encode("utf-8-sig"),
            f"student_{student['seat_no']}_details.csv",
            "text/csv",
        )


# =========================================================
# 14. STUDENT RISK
# =========================================================

elif selected == "Student Risk":
    st.header("⚠️ Students Requiring Attention")

    c1, c2 = st.columns(2)

    with c1:
        percentage_threshold = st.slider(
            "Flag students below percentage",
            min_value=0,
            max_value=100,
            value=40,
        )

    with c2:
        subject_threshold = st.slider(
            "Flag subject marks below",
            min_value=0,
            max_value=100,
            value=40,
        )

    percentages = numeric_series(df["percentage"])
    low_percentage = percentages.lt(
        percentage_threshold
    ).fillna(False)

    low_subject = pd.Series(False, index=df.index)
    risk_reasons = {i: [] for i in df.index}

    for i in df.index:
        if fail_mask.loc[i]:
            risk_reasons[i].append("Official result: FAIL")

        if atkt_mask.loc[i]:
            risk_reasons[i].append("Official result: ATKT")

        if low_percentage.loc[i]:
            risk_reasons[i].append(
                f"Percentage below {percentage_threshold}%"
            )

    for subject in subjects:
        marks = numeric_series(df[subject])
        flagged = marks.lt(subject_threshold).fillna(False)
        low_subject |= flagged

        for i in df.index[flagged]:
            risk_reasons[i].append(
                f"{display_subject(subject, subject_names)}: "
                f"{marks.loc[i]:g} marks"
            )

    risk_mask = fail_mask | atkt_mask | low_percentage | low_subject
    risk_df = df.loc[risk_mask].copy()
    risk_df["Risk Reason"] = [
        "; ".join(risk_reasons[i]) for i in risk_df.index
    ]

    a, b, c = st.columns(3)
    a.metric("Students Flagged", len(risk_df))
    b.metric("Official Fail / ATKT", int((fail_mask | atkt_mask).sum()))
    c.metric("Low Subject Marks", int(low_subject.sum()))

    if risk_df.empty:
        st.success("No students matched the selected thresholds.")
    else:
        visible_columns = [
            col for col in [
                "seat_no", "student_name", "result", "total_marks",
                "percentage", "sgpa", "Risk Reason",
            ]
            if col in risk_df.columns
        ]

        st.dataframe(
            risk_df[visible_columns],
            use_container_width=True,
            hide_index=True,
        )

        st.download_button(
            "Download Student Risk CSV",
            risk_df.to_csv(index=False).encode("utf-8-sig"),
            "student_risk_report.csv",
            "text/csv",
        )

    st.caption(
        "These are rule-based flags, not predictions of a student's future result."
    )


# =========================================================
# 15. STUDENT EXPLORER
# =========================================================

elif selected == "Student Explorer":
    st.header("🔍 Student Explorer")

    query = st.text_input(
        "Search by name or seat number",
        placeholder="Enter student name or seat number",
        key="explorer_search",
    )

    if query.strip():
        matches = df[
            df["student_name"].astype(str).str.contains(
                query, case=False, regex=False, na=False
            )
            | df["seat_no"].astype(str).str.contains(
                query, case=False, regex=False, na=False
            )
        ]
    else:
        matches = df

    st.write(f"Matching students: {len(matches)}")

    if not matches.empty:
        idx = st.selectbox(
            "Select Student",
            matches.index.tolist(),
            format_func=lambda i: (
                f"{df.loc[i, 'student_name']} · {df.loc[i, 'seat_no']}"
            ),
            key="explorer_select",
        )

        row = df.loc[idx]

        a, b, c, d = st.columns(4)

        a.metric("Result", str(row.get("result", "N/A")))

        total = row.get("total_marks", pd.NA)
        b.metric(
            "Total Marks",
            f"{total:g}" if pd.notna(total) else "N/A",
        )

        percentage = row.get("percentage", pd.NA)
        c.metric(
            "Percentage",
            f"{percentage:.2f}%"
            if pd.notna(percentage) else "N/A",
        )

        sgpa = row.get("sgpa", pd.NA)
        d.metric(
            "SGPA",
            f"{sgpa:.2f}" if pd.notna(sgpa) else "N/A",
        )

        if subjects:
            chart_data = pd.DataFrame({
                "Subject": [
                    display_subject(s, subject_names) for s in subjects
                ],
                "Marks": [
                    pd.to_numeric(
                        pd.Series([row.get(s, pd.NA)]),
                        errors="coerce",
                    ).iloc[0]
                    for s in subjects
                ],
            })

            st.plotly_chart(
                px.bar(
                    chart_data.dropna(subset=["Marks"]),
                    x="Subject",
                    y="Marks",
                    title="Subject-wise Marks",
                ),
                use_container_width=True,
            )

            pie_data = chart_data.dropna(subset=["Marks"])
            pie_data = pie_data[pie_data["Marks"] > 0]

            if not pie_data.empty:
                st.plotly_chart(
                    px.pie(
                        pie_data,
                        names="Subject",
                        values="Marks",
                        title="Subject-wise Share of Marks",
                        hole=0.4,
                    ),
                    use_container_width=True,
                )

        with st.expander("All Student Details"):
            st.dataframe(
                pd.DataFrame({
                    "Field": df.columns,
                    "Value": [
                        "" if pd.isna(row.get(col))
                        else str(row.get(col))
                        for col in df.columns
                    ],
                }),
                use_container_width=True,
                hide_index=True,
            )


# =========================================================
# 16. IMPORT RESULTS / DATA QUALITY
# =========================================================

elif selected == "Import Results":
    st.header("📂 Import Results & Data Quality")

    st.write(f"**Source:** {source_name}")
    st.write(f"**Import method:** {import_method}")
    st.write(f"**Semester / Exam:** {semester}")

    st.subheader("Detected Data")
    st.dataframe(
        df.head(200),
        use_container_width=True,
        hide_index=True,
    )

    st.subheader("Detected Columns")
    st.code(", ".join(map(str, df.columns)))

    a, b, c = st.columns(3)

    a.metric("Rows Imported", len(df))
    b.metric(
        "Duplicate Seat Numbers",
        int(df["seat_no"].duplicated().sum()),
    )
    c.metric(
        "Blank Student Names",
        int(df["student_name"].astype(str).str.strip().eq("").sum()),
    )

    if df["seat_no"].duplicated().any():
        st.warning(
            "Duplicate seat numbers were found. Verify these records against the original."
        )

    if df["student_name"].astype(str).str.strip().eq("").any():
        st.warning("Some student names are blank.")

    if import_method.startswith("PDF"):
        st.warning(
            "PDF extraction depends on the document layout. "
            "Verify extracted names, marks and results against the original PDF."
        )


# =========================================================
# 17. EXPORT CENTRE
# =========================================================

elif selected == "Export Centre":
    st.header("📥 Export Centre")

    st.write(
        "Download normalized student data and a detailed Word report."
    )

    left, right = st.columns(2)

    left.download_button(
        "Download CSV",
        df.to_csv(index=False).encode("utf-8-sig"),
        "student_results.csv",
        "text/csv",
        use_container_width=True,
    )

    right.download_button(
        "Download Excel Workbook",
        excel_bytes(df),
        "student_results.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

    try:
        report = word_report(
            df, subjects, subject_names, source_name, semester
        )

        st.download_button(
            "Download Detailed Word Report",
            report,
            "student_result_intelligence_report.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    except Exception as exc:
        st.error(f"Word report could not be generated: {exc}")

    st.caption(
        "Only detected values are included. Missing official marks or statuses "
        "are not invented."
    )
