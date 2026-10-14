
import io
import re
import os
import tempfile
import warnings
from pathlib import Path

import streamlit as st
import pandas as pd
import plotly.express as px
import pdfplumber

warnings.filterwarnings("ignore")

# Optional OCR libraries. OCR also needs Tesseract installed separately.
try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except Exception:
    OCR_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except Exception:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    )
    from reportlab.lib.styles import getSampleStyleSheet
    PDF_EXPORT_AVAILABLE = True
except Exception:
    PDF_EXPORT_AVAILABLE = False


# =========================================================
# PAGE CONFIG + DESIGN
# =========================================================
st.set_page_config(
    page_title="Student Result Intelligence",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
:root {
    --navy: #172554;
    --indigo: #4338ca;
    --indigo-dark: #312e81;
    --bg: #f5f7ff;
    --card: #ffffff;
    --text: #172033;
    --muted: #475569;
    --border: #dbe2f0;
}

.stApp {
    background: var(--bg);
    color: var(--text);
}

[data-testid="stHeader"] {
    background: rgba(245,247,255,0.96);
}

[data-testid="stSidebar"] {
    background: #e9edff;
    border-right: 1px solid #c7d2fe;
}

[data-testid="stSidebar"] * {
    color: #172554 !important;
}

.stApp p, .stApp label, .stApp li,
.stApp h1, .stApp h2, .stApp h3,
.stApp h4, .stApp span,
.stApp [data-testid="stMarkdownContainer"] {
    color: var(--text);
}

h1, h2, h3 {
    color: #172554 !important;
    font-weight: 800 !important;
}

p, label, li {
    color: #273449 !important;
}

.stButton > button,
.stDownloadButton > button {
    background: #4338ca !important;
    color: #ffffff !important;
    border: 1px solid #3730a3 !important;
    border-radius: 9px !important;
    min-height: 42px;
    font-weight: 700 !important;
}

.stButton > button *,
.stDownloadButton > button * {
    color: #ffffff !important;
}

.stButton > button:hover,
.stDownloadButton > button:hover {
    background: #312e81 !important;
    border-color: #1e1b4b !important;
}

.stTextInput input,
.stNumberInput input,
.stSelectbox div[data-baseweb="select"],
.stMultiSelect div[data-baseweb="select"] {
    background: #ffffff !important;
    color: #172033 !important;
    border-color: #a5b4fc !important;
}

div[data-testid="stMetric"] {
    background: #ffffff;
    border: 1px solid #dbe2f0;
    padding: 17px;
    border-radius: 14px;
    box-shadow: 0 2px 8px rgba(30,41,59,0.05);
}

div[data-testid="stMetricLabel"] {
    color: #475569 !important;
}

div[data-testid="stMetricValue"] {
    color: #312e81 !important;
    font-weight: 800;
}

div[data-testid="stDataFrame"] {
    background: white;
    border-radius: 10px;
}

div[data-testid="stAlert"] {
    border-radius: 10px;
}

hr {
    border-color: #dbe2f0;
}

.small-note {
    color: #475569;
    font-size: 0.88rem;
}

.hero {
    background: linear-gradient(120deg, #eef2ff, #ffffff);
    border: 1px solid #c7d2fe;
    border-radius: 18px;
    padding: 22px 25px;
    margin-bottom: 18px;
}
</style>
""", unsafe_allow_html=True)


# =========================================================
# HELPERS
# =========================================================
def normalize_name(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def numeric_series(series):
    return pd.to_numeric(
        series.astype(str).str.replace(",", "", regex=False)
        .str.replace("%", "", regex=False).str.strip(),
        errors="coerce"
    )


def read_uploaded_table(upload):
    name = upload.name.lower()
    upload.seek(0)

    if name.endswith(".csv"):
        try:
            return pd.read_csv(upload, encoding="utf-8-sig")
        except UnicodeDecodeError:
            upload.seek(0)
            return pd.read_csv(upload, encoding="latin-1")

    if name.endswith(".xlsx") or name.endswith(".xls"):
        return pd.read_excel(upload)

    raise ValueError("Unsupported file type. Upload PDF, CSV, XLSX or XLS.")


def canonicalize_columns(df):
    """Map common column labels without inventing absent information."""
    df = df.copy()
    df.columns = [normalize_name(c) for c in df.columns]

    mapping = {}
    for col in df.columns:
        k = normalize_key(col)

        if k in {
            "studentname", "nameofstudent", "candidate", "candidatename",
            "nameofcandidate", "student", "fullname", "name"
        }:
            mapping[col] = "Student Name"
        elif k in {
            "seatno", "seatnumber", "rollno", "rollnumber", "studentid",
            "enrollmentno", "enrolmentno", "prn", "registrationnumber",
            "examnumber"
        }:
            mapping[col] = "Seat Number"
        elif k in {"result", "finalresult", "status", "examresult"}:
            mapping[col] = "Result"
        elif k in {"sgpa", "semestergradepointaverage"}:
            mapping[col] = "SGPA"
        elif k in {"cgpa", "cumulativegradepointaverage"}:
            mapping[col] = "CGPA"
        elif k in {
            "percentage", "percent", "percentageobtained", "percentagemarks",
            "markpercentage", "overallpercentage"
        }:
            mapping[col] = "Percentage"
        elif k in {
            "total", "totalmarks", "markstotal", "marksobtained",
            "totalmarksobtained", "grandtotal", "obtainedmarks"
        }:
            mapping[col] = "Total Marks"
        elif k in {"maxmarks", "maximummarks", "totalmaximummarks"}:
            mapping[col] = "Maximum Marks"
        elif k in {"subject", "subjectname", "coursename", "papername"}:
            mapping[col] = "Subject"
        elif k in {"subjectcode", "coursecode", "papercode"}:
            mapping[col] = "Subject Code"
        elif k in {"teacher", "teachername", "faculty", "facultyname"}:
            mapping[col] = "Teacher"

    df = df.rename(columns=mapping)
    df = df.loc[:, ~df.columns.duplicated()]
    return df


def find_column(df, candidates):
    keys = {normalize_key(c): c for c in df.columns}
    for candidate in candidates:
        if normalize_key(candidate) in keys:
            return keys[normalize_key(candidate)]
    return None


def clean_student_table(df):
    if df is None or df.empty:
        return pd.DataFrame()

    df = canonicalize_columns(df)
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")

    name_col = find_column(df, ["Student Name"])
    seat_col = find_column(df, ["Seat Number"])
    result_col = find_column(df, ["Result"])
    sgpa_col = find_column(df, ["SGPA"])
    percent_col = find_column(df, ["Percentage"])
    total_col = find_column(df, ["Total Marks"])
    max_col = find_column(df, ["Maximum Marks"])

    # Do not treat arbitrary tables as student records.
    if name_col is None and seat_col is None:
        return pd.DataFrame()

    if name_col:
        df["Student Name"] = df[name_col].astype(str).map(normalize_name)
        df.loc[df["Student Name"].str.lower().isin(
            ["", "nan", "none", "student name", "name"]
        ), "Student Name"] = pd.NA

    if seat_col:
        df["Seat Number"] = df[seat_col].astype(str).str.strip()
        df.loc[df["Seat Number"].str.lower().isin(
            ["", "nan", "none", "seat number", "roll no"]
        ), "Seat Number"] = pd.NA

    if result_col:
        df["Result"] = df[result_col].astype(str).str.upper().str.strip()

    for col in ["SGPA", "CGPA", "Percentage", "Total Marks", "Maximum Marks"]:
        if col in df.columns:
            df[col] = numeric_series(df[col])

    # Calculate percentage only when a valid maximum is present.
    if "Percentage" not in df.columns:
        if "Total Marks" in df.columns and "Maximum Marks" in df.columns:
            valid = df["Maximum Marks"].gt(0)
            df["Percentage"] = pd.NA
            df.loc[valid, "Percentage"] = (
                df.loc[valid, "Total Marks"] /
                df.loc[valid, "Maximum Marks"] * 100
            )

    if "Result" in df.columns:
        df["Result"] = df["Result"].replace({
            "NAN": pd.NA, "NONE": pd.NA, "": pd.NA
        })

    # Remove repeated table headers and rows with no identifiable student.
    if "Student Name" in df.columns:
        df = df[
            ~df["Student Name"].astype(str).str.lower().isin(
                ["student name", "name", "nan", "none", ""]
            )
        ]

    if "Seat Number" in df.columns:
        df = df[
            ~df["Seat Number"].astype(str).str.lower().isin(
                ["seat number", "roll number", "nan", "none", ""]
            )
        ]

    df = df.dropna(how="all").reset_index(drop=True)

    # De-duplicate only exact repeated records; do not merge conflicting records.
    df = df.drop_duplicates().reset_index(drop=True)
    return df


def extract_pdf(upload, allow_ocr=False):
    """
    Try text and table extraction page by page.
    OCR is attempted only if enabled and its optional libraries are installed.
    PDF layout differences can still require format-specific rules.
    """
    upload.seek(0)
    pdf_bytes = upload.read()
    page_texts = []
    table_frames = []
    page_info = []

    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                page_texts.append(text)

                try:
                    tables = page.extract_tables()
                except Exception:
                    tables = []

                for table in tables or []:
                    if not table or len(table) < 2:
                        continue

                    rows = []
                    for row in table:
                        if row:
                            rows.append([
                                normalize_name(cell) if cell is not None else ""
                                for cell in row
                            ])

                    if len(rows) >= 2:
                        # Locate a plausible header row in the first few rows.
                        header_idx = None
                        for idx, row in enumerate(rows[:4]):
                            joined = " ".join(row).lower()
                            has_name = any(x in joined for x in [
                                "name", "student", "candidate"
                            ])
                            has_id = any(x in joined for x in [
                                "seat", "roll", "prn", "enrol", "register"
                            ])
                            if has_name or has_id:
                                header_idx = idx
                                break

                        if header_idx is not None:
                            headers = rows[header_idx]
                            width = len(headers)
                            data_rows = [
                                (r + [""] * width)[:width]
                                for r in rows[header_idx + 1:]
                            ]
                            if data_rows:
                                table_frames.append(pd.DataFrame(
                                    data_rows, columns=headers
                                ))

                page_info.append({
                    "Page": page_num,
                    "Text characters": len(text),
                    "Tables found": len(tables or [])
                })

    except Exception as exc:
        raise ValueError(
            "PDF open karta ala nahi. File damaged, encrypted kiwa unsupported asu shakate. "
            f"Technical detail: {exc}"
        ) from exc

    extracted = None
    method = None

    # First choice: PDF tables.
    for candidate in table_frames:
        cleaned = clean_student_table(candidate)
        if not cleaned.empty and (
            "Student Name" in cleaned.columns or
            "Seat Number" in cleaned.columns
        ):
            if extracted is None or len(cleaned) > len(extracted):
                extracted = cleaned
                method = "PDF table extraction"

    # Second choice: try common DBATU-style text records.
    if extracted is None or extracted.empty:
        records = []
        full_text = "\n".join(page_texts)

        # The student ID and PASS/FAIL act as anchors. The remaining fields
        # are kept only when they can be recognized with reasonable confidence.
        pattern = re.compile(
            r"(?P<seat>\b\d{10,16}\b)\s+"
            r"(?P<name>[A-Z][A-Z .'-]{2,}?)\s+"
            r"(?P<result>PASS|FAIL|ATKT|A\.T\.K\.T\.|ABSENT)\b",
            flags=re.IGNORECASE
        )

        for match in pattern.finditer(full_text):
            records.append({
                "Seat Number": match.group("seat"),
                "Student Name": normalize_name(match.group("name")),
                "Result": match.group("result").upper().replace(".", "")
            })

        if records:
            extracted = clean_student_table(pd.DataFrame(records))
            method = "PDF text pattern extraction"

    # Third choice: OCR for scanned pages, if explicitly enabled.
    if (extracted is None or extracted.empty) and allow_ocr:
        if not OCR_AVAILABLE:
            raise ValueError(
                "OCR enable kele aahe pan pytesseract/Pillow install nahi. "
                "OCR dependencies install kara ani Tesseract OCR application pan install kara."
            )

        try:
            import fitz
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            ocr_texts = []

            for page in doc:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_texts.append(pytesseract.image_to_string(image))

            full_ocr = "\n".join(ocr_texts)
            records = []
            for match in pattern.finditer(full_ocr):
                records.append({
                    "Seat Number": match.group("seat"),
                    "Student Name": normalize_name(match.group("name")),
                    "Result": match.group("result").upper().replace(".", "")
                })

            if records:
                extracted = clean_student_table(pd.DataFrame(records))
                method = "OCR text pattern extraction"

        except Exception as exc:
            raise ValueError(
                "Scanned PDF cha OCR complete zala nahi. Tesseract installation "
                f"check kara. Technical detail: {exc}"
            ) from exc

    if extracted is None or extracted.empty:
        raise ValueError(
            "PDF madhun reliable student table sapadla nahi. Ha PDF scanned asu shakto "
            "kinva tyacha layout vegla asu shakto. App ne marks andazane bharlele nahit. "
            "OCR enable kara, kiwa PDF madhil result table Excel/CSV madhye export karun upload kara."
        )

    return extracted, method, pd.DataFrame(page_info), page_texts


def student_quality_checks(df):
    checks = []
    name_col = find_column(df, ["Student Name"])
    seat_col = find_column(df, ["Seat Number"])

    checks.append({
        "Validation": "Student records extracted",
        "Status": "PASS" if len(df) > 0 else "FAIL",
        "Details": f"{len(df)} rows"
    })

    if name_col:
        valid_names = df[name_col].notna().sum()
        checks.append({
            "Validation": "Student names",
            "Status": "PASS" if valid_names > 0 else "WARNING",
            "Details": f"{valid_names}/{len(df)} rows have names"
        })

    if seat_col:
        valid_ids = df[seat_col].notna().sum()
        duplicates = df[seat_col].dropna().duplicated().sum()
        checks.append({
            "Validation": "Seat numbers",
            "Status": "PASS" if valid_ids > 0 and duplicates == 0 else "WARNING",
            "Details": f"{valid_ids}/{len(df)} populated; {duplicates} duplicate IDs"
        })

    if "Percentage" in df.columns:
        valid_pct = df["Percentage"].dropna()
        in_range = valid_pct.between(0, 100).all()
        checks.append({
            "Validation": "Percentage range",
            "Status": "PASS" if in_range else "WARNING",
            "Details": "Values are within 0–100" if in_range else
                       "Some percentages are outside 0–100; review source data"
        })

    if "SGPA" in df.columns:
        valid_sgpa = df["SGPA"].dropna()
        plausible = valid_sgpa.between(0, 10).all()
        checks.append({
            "Validation": "SGPA range",
            "Status": "PASS" if plausible else "WARNING",
            "Details": "Values are within 0–10" if plausible else
                       "Some SGPA values are outside 0–10; review source data"
        })

    return pd.DataFrame(checks)


def make_excel(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Student Results")
        student_quality_checks(df).to_excel(
            writer, index=False, sheet_name="Validation"
        )
    output.seek(0)
    return output.getvalue()


def make_word(df, title="Student Result Analysis Report"):
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed.")

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Total extracted records: {len(df)}")
    doc.add_paragraph(
        "This report uses only fields found in the uploaded file. "
        "Missing teacher names or subject marks are not inferred."
    )

    doc.add_heading("Validation", level=1)
    for _, row in student_quality_checks(df).iterrows():
        doc.add_paragraph(
            f"{row['Validation']}: {row['Status']} — {row['Details']}"
        )

    doc.add_heading("Student Records", level=1)
    preview = df.head(150)
    table = doc.add_table(rows=1, cols=len(preview.columns))
    table.style = "Table Grid"

    for i, col in enumerate(preview.columns):
        table.rows[0].cells[i].text = str(col)

    for _, row in preview.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if pd.isna(value) else str(value)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.getvalue()


def make_pdf_report(df):
    if not PDF_EXPORT_AVAILABLE:
        raise RuntimeError("reportlab is not installed.")

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4),
        rightMargin=24, leftMargin=24,
        topMargin=24, bottomMargin=24
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Student Result Analysis Report", styles["Title"]),
        Spacer(1, 10),
        Paragraph(f"Extracted student records: {len(df)}", styles["Normal"]),
        Spacer(1, 12)
    ]

    report_df = df.head(150).fillna("")
    data = [list(map(str, report_df.columns))]
    for row in report_df.itertuples(index=False, name=None):
        data.append([str(v)[:60] for v in row])

    # Keep very wide datasets readable by limiting columns in the PDF.
    max_cols = 8
    if len(data[0]) > max_cols:
        keep = list(range(max_cols))
        data = [[row[i] for i in keep] for row in data]

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312e81")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#eef2ff")])
    ]))
    story.append(table)
    doc.build(story)
    out.seek(0)
    return out.getvalue()


# =========================================================
# HEADER
# =========================================================
st.markdown("""
<div class="hero">
    <h1>🎓 Student Result Intelligence</h1>
    <p>Academic performance analytics, student lookup, result validation and reports.</p>
</div>
""", unsafe_allow_html=True)

st.caption(
    "Upload a result PDF, CSV or Excel workbook. Check extracted records before relying on analysis."
)


# =========================================================
# SIDEBAR: UPLOAD AND NAVIGATION
# =========================================================
with st.sidebar:
    st.markdown("## 🎓 Student Intelligence")
    st.markdown("---")
    uploaded_file = st.file_uploader(
        "Upload result file",
        type=["pdf", "csv", "xlsx", "xls"],
        help="Text PDFs and table PDFs are supported. Scanned PDFs may need OCR."
    )
    enable_ocr = st.checkbox(
        "Try OCR for scanned PDFs",
        value=False,
        help="Requires pytesseract, Pillow, PyMuPDF and the Tesseract application."
    )

    st.markdown("---")
    st.markdown("### Navigation")
    page = st.radio(
        "Select a section",
        [
            "Overview",
            "Student Explorer",
            "Top Performers",
            "Student Risk",
            "Subject Analysis",
            "All Records",
            "Validation & Export"
        ],
        index=0
    )

    st.markdown("---")
    st.markdown(
        "<p class='small-note'>The app does not invent missing marks, teacher names or results.</p>",
        unsafe_allow_html=True
    )


# =========================================================
# FILE LOADING
# =========================================================
if uploaded_file is None:
    st.info("👈 Start by uploading a student result PDF, CSV or Excel file.")
    st.markdown("### What you can analyse")
    a, b, c = st.columns(3)
    a.metric("Student lookup", "Available")
    b.metric("Performance charts", "Available")
    c.metric("Export formats", "Excel / Word / PDF")
    st.stop()

try:
    with st.spinner("Reading file and extracting available records..."):
        if uploaded_file.name.lower().endswith(".pdf"):
            df, extraction_method, page_details, extracted_text = extract_pdf(
                uploaded_file, allow_ocr=enable_ocr
            )
        else:
            df = clean_student_table(read_uploaded_table(uploaded_file))
            extraction_method = "CSV/Excel import"
            page_details = pd.DataFrame()
            extracted_text = []

    if df.empty:
        st.error(
            "File read zali, pan student records identify zhale nahit. "
            "Column names madhye Student Name, Seat Number, Result, Percentage "
            "kinva Total Marks asne upyogi aahe."
        )
        st.stop()

except Exception as exc:
    st.error("File analysis complete zali nahi.")
    st.warning(str(exc))
    st.markdown("### Troubleshooting")
    st.markdown("""
    1. PDF password-protected naslyachi khatri kara.
    2. Scanned PDF asel tar OCR option enable kara.
    3. OCR sathi Tesseract application system var install asne garjeche aahe.
    4. PDF madhil table Excel/CSV madhye export karta yet asel tar te upload kara.
    5. Marks cha maximum mahit nasel tar app percentage andazane calculate karat nahi.
    """)
    st.stop()


# =========================================================
# DATA PREPARATION
# =========================================================
name_col = find_column(df, ["Student Name"])
seat_col = find_column(df, ["Seat Number"])
result_col = find_column(df, ["Result"])
percent_col = find_column(df, ["Percentage"])
sgpa_col = find_column(df, ["SGPA"])
total_col = find_column(df, ["Total Marks"])
subject_col = find_column(df, ["Subject"])
subject_code_col = find_column(df, ["Subject Code"])
teacher_col = find_column(df, ["Teacher"])

st.success(
    f"File read successful. Method: {extraction_method}. "
    f"Extracted records: {len(df)}"
)

quality = student_quality_checks(df)
has_numeric_performance = (
    (percent_col is not None and df[percent_col].notna().any())
    or (sgpa_col is not None and df[sgpa_col].notna().any())
    or (total_col is not None and df[total_col].notna().any())
)

# If the file contains multiple rows per student (e.g. subject marks),
# do not pretend each row is one unique student.
if seat_col:
    unique_students = df[seat_col].nunique(dropna=True)
elif name_col:
    unique_students = df[name_col].nunique(dropna=True)
else:
    unique_students = len(df)

with st.expander("Extraction details and validation", expanded=True):
    st.write(f"**Extraction method:** {extraction_method}")
    st.write(f"**Rows extracted:** {len(df)}")
    st.write(f"**Distinct student IDs/names:** {unique_students}")
    st.dataframe(quality, use_container_width=True, hide_index=True)

    if page_details is not None and not page_details.empty:
        st.caption("PDF page-level extraction details")
        st.dataframe(page_details, use_container_width=True, hide_index=True)

    if name_col is None and seat_col is None:
        st.error(
            "No student name or ID column was detected. Review the uploaded table "
            "before using student-level analytics."
        )
        st.stop()

    if len(df) < 2:
        st.warning(
            "Only one record was extracted. Verify the source PDF/table before "
            "treating this as a complete class result."
        )

    if seat_col and df[seat_col].dropna().duplicated().any():
        st.warning(
            "Repeated seat numbers found. This can happen in subject-wise data "
            "or from duplicate extraction. Review All Records before exporting."
        )

if not has_numeric_performance:
    st.warning(
        "Student identity/result fields were found, but numeric marks, percentage "
        "or SGPA were not detected. Student lookup is available; numeric performance "
        "charts and rankings need valid numeric marks."
    )


# =========================================================
# OVERVIEW
# =========================================================
if page == "Overview":
    st.subheader("📊 Class Overview")

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Extracted rows", f"{len(df):,}")
    m2.metric("Distinct students", f"{unique_students:,}")

    if result_col:
        result_values = df[result_col].astype(str).str.upper()
        pass_count = result_values.isin(["PASS", "PASSED"]).sum()
        fail_count = result_values.isin(["FAIL", "FAILED"]).sum()
        m3.metric("Pass records", f"{pass_count:,}")
        m4.metric("Fail records", f"{fail_count:,}")
    elif percent_col:
        m3.metric("Average percentage", (
            f"{df[percent_col].mean():.2f}%"
            if df[percent_col].notna().any() else "N/A"
        ))
        m4.metric("Highest percentage", (
            f"{df[percent_col].max():.2f}%"
            if df[percent_col].notna().any() else "N/A"
        ))
    elif sgpa_col:
        m3.metric("Average SGPA", (
            f"{df[sgpa_col].mean():.2f}"
            if df[sgpa_col].notna().any() else "N/A"
        ))
        m4.metric("Highest SGPA", (
            f"{df[sgpa_col].max():.2f}"
            if df[sgpa_col].notna().any() else "N/A"
        ))
    else:
        m3.metric("Numeric fields", "Not detected")
        m4.metric("Analysis status", "Limited")

    left, right = st.columns(2)

    with left:
        if percent_col and df[percent_col].notna().any():
            fig = px.histogram(
                df, x=percent_col, nbins=12,
                title="Percentage Distribution",
                color_discrete_sequence=["#4338ca"]
            )
            fig.update_layout(
                paper_bgcolor="white", plot_bgcolor="white",
                font_color="#172033"
            )
            st.plotly_chart(fig, use_container_width=True)
        elif sgpa_col and df[sgpa_col].notna().any():
            fig = px.histogram(
                df, x=sgpa_col, nbins=10,
                title="SGPA Distribution",
                color_discrete_sequence=["#4338ca"]
            )
            fig.update_layout(
                paper_bgcolor="white", plot_bgcolor="white",
                font_color="#172033"
            )
            st.plotly_chart(fig, use_container_width=True)
        else:
            st.info("Percentage/SGPA distribution sathi numeric column available nahi.")

    with right:
        if result_col:
            counts = (
                df[result_col].astype(str).str.upper()
                .value_counts().rename_axis("Result").reset_index(name="Count")
            )
            fig = px.pie(
                counts, names="Result", values="Count",
                title="Result Status Distribution",
                color_discrete_sequence=["#4338ca", "#ef4444", "#f59e0b", "#0f766e"]
            )
            fig.update_layout(
                paper_bgcolor="white", font_color="#172033"
            )
            st.plotly_chart(fig, use_container_width=True)
        elif total_col and df[total_col].notna().any():
            fig = px.box(
                df, y=total_col, title="Total Marks Distribution",
                color_discrete_sequence=["#4338ca"]
            )
            fig.update_layout(
                paper_bgcolor="white", plot_bgcolor="white",
                font_color="#172033"
            )
            st.plotly_chart(fig, use_container_width=True)
        else:
            st.info("Result status/total marks column available nahi.")

    st.subheader("Data Preview")
    st.dataframe(df.head(20), use_container_width=True, hide_index=True)


# =========================================================
# STUDENT EXPLORER
# =========================================================
elif page == "Student Explorer":
    st.subheader("🔎 Student Explorer")

    query = st.text_input(
        "Search by student name or seat/roll number",
        placeholder="Enter name or seat number..."
    )

    mask = pd.Series(True, index=df.index)
    if query.strip():
        query = query.strip().lower()
        mask = pd.Series(False, index=df.index)
        if name_col:
            mask |= df[name_col].astype(str).str.lower().str.contains(
                re.escape(query), na=False
            )
        if seat_col:
            mask |= df[seat_col].astype(str).str.lower().str.contains(
                re.escape(query), na=False
            )

    matches = df[mask]
    st.write(f"Matching records: **{len(matches)}**")

    if matches.empty:
        st.warning("Matching student sapadla nahi.")
    else:
        st.dataframe(matches, use_container_width=True, hide_index=True)

        if len(matches) == 1:
            row = matches.iloc[0]
            st.subheader("Student Details")
            st.json({
                str(k): (None if pd.isna(v) else str(v))
                for k, v in row.to_dict().items()
            })

            chart_col = percent_col or sgpa_col or total_col
            if chart_col and pd.notna(row[chart_col]):
                st.metric(chart_col, f"{row[chart_col]:,.2f}")


# =========================================================
# TOP PERFORMERS
# =========================================================
elif page == "Top Performers":
    st.subheader("🏆 Top Performers")

    ranking_col = percent_col or sgpa_col or total_col

    if not ranking_col or not df[ranking_col].notna().any():
        st.warning(
            "Ranking sathi valid Percentage, SGPA kiwa Total Marks column pahije."
        )
    else:
        top_n = st.slider("Number of records to show", 5, 50, 10)
        ranked = df.dropna(subset=[ranking_col]).sort_values(
            ranking_col, ascending=False
        ).head(top_n)
        st.dataframe(ranked, use_container_width=True, hide_index=True)

        if name_col and ranked[name_col].notna().any():
            chart = ranked.copy()
            chart["Chart Label"] = chart[name_col].astype(str)
            if seat_col:
                chart["Chart Label"] = (
                    chart["Chart Label"] + " (" +
                    chart[seat_col].astype(str) + ")"
                )
            fig = px.bar(
                chart, x="Chart Label", y=ranking_col,
                title=f"Top {top_n} by {ranking_col}",
                color_discrete_sequence=["#4338ca"]
            )
            fig.update_layout(
                xaxis_tickangle=-35, paper_bgcolor="white",
                plot_bgcolor="white", font_color="#172033"
            )
            st.plotly_chart(fig, use_container_width=True)


# =========================================================
# STUDENT RISK
# =========================================================
elif page == "Student Risk":
    st.subheader("⚠️ Student Risk Review")

    if percent_col and df[percent_col].notna().any():
        threshold = st.slider("Low-percentage threshold", 0, 100, 40)
        risk_df = df[
            df[percent_col].notna() &
            (df[percent_col] < threshold)
        ].sort_values(percent_col)
        st.write(f"Records below {threshold}%: **{len(risk_df)}**")
        st.dataframe(risk_df, use_container_width=True, hide_index=True)

    elif result_col:
        failed = df[
            df[result_col].astype(str).str.upper().isin(
                ["FAIL", "FAILED", "ATKT", "A.T.K.T"]
            )
        ]
        st.write(f"Fail/ATKT records found: **{len(failed)}**")
        st.dataframe(failed, use_container_width=True, hide_index=True)

    elif sgpa_col and df[sgpa_col].notna().any():
        threshold = st.slider("Low-SGPA threshold", 0.0, 10.0, 5.0, 0.5)
        risk_df = df[df[sgpa_col].notna() & (df[sgpa_col] < threshold)]
        st.write(f"Records below SGPA {threshold}: **{len(risk_df)}**")
        st.dataframe(risk_df, use_container_width=True, hide_index=True)

    else:
        st.info(
            "Risk review sathi Percentage, SGPA kiwa Result column avashyak aahe. "
            "Missing values la fail samajle jat nahi."
        )


# =========================================================
# SUBJECT ANALYSIS
# =========================================================
elif page == "Subject Analysis":
    st.subheader("📚 Subject Analysis")

    if subject_col is None:
        st.warning(
            "Uploaded data madhye Subject/Subject Name column sapadla nahi. "
            "Student-wise result PDF madhye subject marks separate rows/columns "
            "madhye asu shaktat. Subject-wise analysis sathi subject-level data upload kara."
        )
    else:
        st.dataframe(
            df[[c for c in [
                subject_code_col, subject_col, teacher_col,
                percent_col, total_col, result_col
            ] if c is not None]].head(500),
            use_container_width=True, hide_index=True
        )

        metric_col = percent_col or total_col
        if metric_col:
            numeric_df = df.copy()
            numeric_df[metric_col] = numeric_series(numeric_df[metric_col])
            summary = numeric_df.groupby(subject_col, dropna=False)[metric_col].agg(
                Average="mean", Highest="max", Lowest="min", Records="count"
            ).reset_index()
            summary = summary.sort_values("Average", ascending=False)
            st.subheader("Subject Summary")
            st.dataframe(summary.round(2), use_container_width=True, hide_index=True)

            fig = px.bar(
                summary, x=subject_col, y="Average",
                title=f"Average {metric_col} by Subject",
                color_discrete_sequence=["#4338ca"]
            )
            fig.update_layout(
                xaxis_tickangle=-30, paper_bgcolor="white",
                plot_bgcolor="white", font_color="#172033"
            )
            st.plotly_chart(fig, use_container_width=True)

        if teacher_col:
            st.caption("Teacher/faculty information was present in the uploaded data.")
        else:
            st.info("Teacher names source file madhye nahit; app te invent karat nahi.")


# =========================================================
# ALL RECORDS
# =========================================================
elif page == "All Records":
    st.subheader("📋 All Extracted Records")

    search = st.text_input("Filter records", placeholder="Type a name, ID, result...")
    view_df = df.copy()

    if search.strip():
        text_mask = view_df.astype(str).apply(
            lambda col: col.str.contains(re.escape(search.strip()), case=False, na=False)
        ).any(axis=1)
        view_df = view_df[text_mask]

    st.write(f"Showing **{len(view_df)}** rows")
    st.dataframe(view_df, use_container_width=True, hide_index=True)

    st.download_button(
        "⬇️ Download filtered CSV",
        data=view_df.to_csv(index=False).encode("utf-8-sig"),
        file_name="student_result_records.csv",
        mime="text/csv"
    )


# =========================================================
# VALIDATION + EXPORT
# =========================================================
elif page == "Validation & Export":
    st.subheader("✅ Validation & Export")

    st.markdown("### Validation results")
    st.dataframe(quality, use_container_width=True, hide_index=True)

    st.markdown("### Export analysed records")
    st.caption(
        "Reports include extracted data only. Review the validation results and "
        "verify student names/marks against the source before official use."
    )

    st.download_button(
        "⬇️ Download Excel report",
        data=make_excel(df),
        file_name="student_result_analysis.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    if DOCX_AVAILABLE:
        try:
            word_bytes = make_word(df)
            st.download_button(
                "⬇️ Download Word report",
                data=word_bytes,
                file_name="student_result_analysis.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        except Exception as exc:
            st.error(f"Word report create karta ala nahi: {exc}")
    else:
        st.warning("Word export sathi python-docx install kara.")

    if PDF_EXPORT_AVAILABLE:
        try:
            pdf_bytes = make_pdf_report(df)
            st.download_button(
                "⬇️ Download PDF report",
                data=pdf_bytes,
                file_name="student_result_analysis.pdf",
                mime="application/pdf"
            )
        except Exception as exc:
            st.error(f"PDF report create karta ala nahi: {exc}")
    else:
        st.warning("PDF export sathi reportlab install kara.")


st.markdown("---")
st.caption(
    "Student Result Intelligence • Verify extracted results against the original "
    "university document before official decisions."
)
